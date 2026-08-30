"""Parse Bill-of-Materials files from suppliers into normalized part rows."""

import csv
import io
import os
import re

# Category recognition lives in categories.py, next to the canonical names it
# produces. Re-exported here because BOM parsing is its main caller.
from categories import guess_category, category_from_footprint  # noqa: F401

HEADER_ALIASES = {
    "quantity": [
        "quantity", "qty", "order qty", "order quantity", "qty ordered",
        "quantity ordered", "qnty", "amount",
    ],
    "part_number": [
        "manufacture part number", "manufacturer part number", "mfr part #",
        "mfr part number", "manufacturer part #", "mpn",
        "manufacturer partno", "mfg part #", "mfg part number",
    ],
    "supplier_pn": [
        "lcsc part number", "lcsc part", "lcsc", "digikey part number",
        "digi-key part number", "digikey part #", "supplier part number",
        "supplier part", "order code", "customer no.", "stock code",
    ],
    "manufacturer": ["manufacturer", "manufacture", "mfr", "brand", "mfg"],
    "description": ["description", "desc", "details", "comment", "specification"],
    "package": ["package", "footprint", "case", "package/case", "case/package"],
    "value": ["value", "comment", "comp value", "val", "designation"],
    "unit_cost": [
        "unit price", "unit price($)", "unit price(usd)", "price", "unit cost",
        "unit price (usd)", "unitprice",
    ],
    "datasheet_url": ["datasheet", "datasheet url", "datasheet link"],
    "designator": [
        "designator", "designators", "reference", "references", "refdes",
        "ref des", "ref", "customer reference",
    ],
}

def _norm_header(h):
    return re.sub(r"\s+", " ", (h or "").strip().lower())


def normalize_kicad_footprint(fp):
    """Extract a common package name from a KiCad footprint string.
    E.g. 'Resistor_SMD:R_0603_1608Metric' -> '0603'."""
    if not fp:
        return ""
    if ":" in fp:
        fp = fp.split(":", 1)[1]
    m = re.match(r"[RCL]_(\d{4})_", fp)
    if m:
        return m.group(1)
    m = re.search(
        r"\b(SOT-?\d+(?:-\d+)?|SO(?:IC)?-\d+|[TLD]?QFP-?\d+|[DQ]FN-?\d+|"
        r"BGA-?\d+|[TM]?S?SOP-?\d+|DIP-?\d+|PDIP-?\d+|TO-?\d+|SC-?\d+|"
        r"SOD-?\d+|SMA|SMB|SMC|DO-?\d+|WLCSP-?\d+|LGA-?\d+)\b",
        fp, re.IGNORECASE,
    )
    if m:
        return m.group(1).upper()
    return fp.split("_")[0] if "_" in fp else fp


def detect_supplier(headers, sample_codes=None):
    joined = " | ".join(_norm_header(h) for h in headers)
    if "lcsc part number" in joined or "lcsc" in joined:
        return "LCSC"
    if "digikey" in joined or "digi-key" in joined:
        return "DigiKey"
    for code in (sample_codes or []):
        c = (code or "").strip().upper()
        if c.endswith("-ND") or c.endswith("-CT") or c.endswith("-TR"):
            return "DigiKey"
        if re.fullmatch(r"C\d{3,}", c):
            return "LCSC"
    return "Generic"


def _build_column_map(headers):
    colmap = {}
    bare_part_idx = None
    for idx, h in enumerate(headers):
        nh = _norm_header(h)
        if nh == "part number":
            bare_part_idx = idx
            continue
        for field, aliases in HEADER_ALIASES.items():
            if nh in aliases and field not in colmap:
                colmap[field] = idx
    if bare_part_idx is not None:
        if "part_number" not in colmap:
            colmap["part_number"] = bare_part_idx
        elif "supplier_pn" not in colmap:
            colmap["supplier_pn"] = bare_part_idx
    return colmap


def _clean_price(raw):
    if raw is None:
        return None
    s = re.sub(r"[^0-9.]+", "", str(raw))
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        return _read_xlsx(path)
    return _read_csv(path)


def _detect_delimiter(text):
    first_line = text.splitlines()[0] if text else ""
    candidates = [",", ";", "\t"]
    counts = {d: first_line.count(d) for d in candidates}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ","


def _read_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        text = f.read()
    delimiter = _detect_delimiter(text)
    reader = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    reader = [r for r in reader if any(c.strip() for c in r)]
    header_idx = _find_header_row(reader)
    headers = reader[header_idx]
    return headers, reader[header_idx + 1:]


def _read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(["" if c is None else str(c) for c in r])
    wb.close()
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    header_idx = _find_header_row(rows)
    headers = rows[header_idx]
    return headers, rows[header_idx + 1:]


def _find_header_row(rows):
    known = {a for aliases in HEADER_ALIASES.values() for a in aliases}
    for i, row in enumerate(rows[:10]):
        if sum(1 for c in row if _norm_header(c) in known) >= 2:
            return i
    return 0


def parse_bom(path):
    headers, raw_rows = _read_rows(path)
    colmap = _build_column_map(headers)
    warnings = []
    if "quantity" not in colmap:
        warnings.append("No quantity column found - defaulting each line to 1.")
    if "part_number" not in colmap and "supplier_pn" not in colmap:
        warnings.append("No part-number column found - rows may be hard to match.")

    def cell(row, field):
        i = colmap.get(field)
        if i is None or i >= len(row):
            return ""
        return str(row[i]).strip()

    rows = []
    for row in raw_rows:
        if not any(str(c).strip() for c in row):
            continue
        qty_raw = cell(row, "quantity")
        try:
            qty = int(float(qty_raw)) if qty_raw else 1
        except ValueError:
            qty = 1
        part = {
            "part_number": cell(row, "part_number"),
            "supplier_pn": cell(row, "supplier_pn"),
            "manufacturer": cell(row, "manufacturer"),
            "description": cell(row, "description"),
            "package": cell(row, "package"),
            "value": cell(row, "value") or cell(row, "description"),
            "quantity": qty,
            "unit_cost": _clean_price(cell(row, "unit_cost")),
            "designator": cell(row, "designator"),
        }
        if not (part["part_number"] or part["supplier_pn"] or part["value"]):
            continue
        part["category"] = guess_category(
            part["description"], part["value"], part["part_number"],
            package=part.get("package"),
        )
        rows.append(part)

    supplier = detect_supplier(headers, [r["supplier_pn"] for r in rows])
    for r in rows:
        r["supplier"] = supplier if supplier != "Generic" else ""

    return {
        "supplier": supplier,
        "columns": sorted(colmap.keys()),
        "rows": rows,
        "warnings": warnings,
    }
